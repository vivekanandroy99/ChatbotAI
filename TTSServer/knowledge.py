"""
Per-avatar document knowledge base. Each avatar owns a folder under
KNOWLEDGE_ROOT (Assets/StreamingAssets/Knowledge/<avatarId>/). Any supported
document dropped in there is extracted, chunked and embedded; the index
hot-reloads whenever files are added, changed or removed.

Embeddings are multilingual (bge-m3), so a Hindi question still matches
English documents and vice versa - and Hindi documents work too: an English
question scored the same against a Hindi document as against its English version
(0.68 vs 0.72, off-topic questions 0.27-0.36). Hindi sentences end in the danda (।).
Hindi PDFs are the weak spot: many store Devanagari in a way text extraction
scrambles (or in old non-Unicode fonts like Kruti Dev, which come out as Latin
gibberish); a warning is logged when extracted Hindi looks broken - use .docx/.txt.
"""
import hashlib
import logging
import re
import threading
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import torch
from transformers import AutoModel, AutoModelForSequenceClassification, AutoTokenizer

from hinglish import restore_english

_DEVANAGARI = re.compile(r"[ऀ-ॿ]")

log = logging.getLogger("knowledge")

EMBED_MODEL_ID = "BAAI/bge-m3"
# Optional second pass (2026-10-07): bge-m3 picks the RERANK_CANDIDATES closest passages, the reranker reads the
# question together with each one and reorders them; the best top_k go to the brain. The relevance gate's scores
# stay bge-m3's, so on/off-topic decisions don't change - only which passages the answer is written from.
RERANK_MODEL_ID = "BAAI/bge-reranker-v2-m3"
RERANK_CANDIDATES = 12
SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".pptx"}
LEGACY = {".doc", ".ppt"}
CHUNK_CHARS = 800
CHUNK_OVERLAP = 150
# Slides/pages shorter than this are merged into the next one - a four-word cover
# slide otherwise out-ranks real content for any question mentioning the brand.
MIN_UNIT_CHARS = 300
# Bump when what gets embedded changes, so cached indexes are rebuilt.
INDEX_FORMAT = 6
# Sentences shorter than this aren't scored on their own (headings, fragments).
MIN_SENTENCE_CHARS = 25
# Leading dot: Unity ignores the folder, so no .meta files get generated in it.
INDEX_DIRNAME = ".index"
AVATAR_NAME = re.compile(r"^[A-Za-z0-9_\- ]+$")
_STOPWORDS = {"the", "and", "of", "for", "to", "in", "on", "a", "an", "by", "with", "final", "copy", "new", "v1", "v2"}


@dataclass
class _Index:
    signature: tuple
    texts: list
    sources: list
    emb: np.ndarray
    brands: dict
    # Every passage's sentences, scored separately: one small fact ("pets are
    # welcome") is drowned out in a passage-level match.
    sent_emb: np.ndarray
    sent_parent: np.ndarray


def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


_RECORD_ID = re.compile(r"^[A-Z]{1,3}\d{2,4}$")


def _useful_columns(rows):
    """Drops table columns that tell the reader nothing: the same value on every row ("Project page", "Logo
    source") or record IDs (P037, C104) - the answer agent copied IDs into replies."""
    width = max((len(r) for r in rows), default=0)
    if len(rows) < 4 or any(len(r) != width for r in rows):
        return rows
    body = rows[1:]
    drop = {c for c in range(width)
            if len({r[c] for r in body}) == 1 or all(_RECORD_ID.match(r[c]) for r in body)}
    if not drop or len(drop) == width:
        return rows
    return [[cell for c, cell in enumerate(r) if c not in drop] for r in rows]


def _extract_units(path: Path):
    """Returns [(source_label, text)] - one unit per page/slide where that exists."""
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return [(path.name, path.read_text(encoding="utf-8", errors="ignore"))]

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [(f"{path.name} p.{i + 1}", page.extract_text() or "") for i, page in enumerate(reader.pages)]

    if ext == ".docx":
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        # Paragraphs and tables in document order, one unit per heading section (like a PDF page), so a section
        # isn't cut mid-way and glued to the next. (Tables used to be appended at the end, away from their heading,
        # and the whole document chunked as one text: the P&I museum list never reached the top passages, 2026-10-07.)
        # A long table repeats its heading every few rows, so each chunk of it still says what it lists.
        document = docx.Document(str(path))
        units, parts, heading = [], [], ""

        def flush():
            if any(t.strip() for t in parts):
                units.append((f"{path.name} - {heading}" if heading else path.name, "\n".join(parts)))
            parts.clear()

        for block in document.element.body.iterchildren():
            tag = block.tag.rsplit("}", 1)[-1]
            if tag == "p":
                p = Paragraph(block, document)
                style = p.style.name.lower() if p.style is not None else ""
                if style.startswith(("heading", "title")) and p.text.strip():
                    flush()
                    heading = p.text.strip()
                parts.append(p.text)
            elif tag == "tbl":
                rows = []
                for row in Table(block, document).rows:
                    cells = []
                    for cell in row.cells:  # merged cells repeat their text
                        if not cells or cell.text != cells[-1]:
                            cells.append(cell.text.strip())
                    rows.append(cells)
                for i, cells in enumerate(_useful_columns(rows)):
                    if i and i % 8 == 0 and heading:
                        parts.append(f"({heading}, continued)")
                    parts.append(" | ".join(cells))
        flush()
        return units

    if ext == ".pptx":
        from pptx import Presentation

        units = []
        for i, slide in enumerate(Presentation(str(path)).slides):
            parts = []
            for shape in _iter_shapes(slide.shapes):
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text for cell in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                parts.append(slide.notes_slide.notes_text_frame.text)
            units.append((f"{path.name} slide {i + 1}", "\n".join(parts)))
        return units

    return []


def _merge_small_units(units):
    merged, pending_label, pending = [], None, ""
    for label, text in units:
        pending_label = pending_label or label
        pending = f"{pending}\n{text}" if pending else text
        if len(pending.strip()) >= MIN_UNIT_CHARS:
            merged.append((pending_label, pending))
            pending_label, pending = None, ""
    if pending.strip():
        if merged:
            merged[-1] = (merged[-1][0], merged[-1][1] + "\n" + pending)
        else:
            merged.append((pending_label, pending))
    return merged


def _title(path: Path) -> str:
    return re.sub(r"[_\-]+", " ", path.stem).strip()


_EMAIL_OR_URL = re.compile(r"\S+@\S+|\b(?:https?://|www\.)\S+|\b[\w\-]+\.(?:com|co|in|org|net|io|ai|app|example)\b\S*", re.I)


def _brand_terms(texts):
    """Words the documents treat as names, lowercased -> their usual spelling.
    A name is written with a trademark sign at least 3 times (altcore ™ -
    stylized all-lowercase brands are why capitalization alone isn't enough),
    in camelCase (AltCore), or capitalized nearly every time it appears
    mid-sentence in real prose (acronyms like PWA fall out of this too).
    Only full sentences count for capitalization: slide decks are full of
    Title Case and ALL CAPS headings."""
    trademarked, marked, spellings = Counter(), set(), Counter()
    seen, capitalized = Counter(), Counter()
    for text in texts:
        # Email addresses and websites write the brand in lowercase (hello@brand.com);
        # counted as words they'd make the brand look like an ordinary word.
        text = _EMAIL_OR_URL.sub(" ", text)
        for m in re.finditer(r"([A-Za-z][A-Za-z0-9\-]+)\s*[™®]", text):
            trademarked[m.group(1).lower()] += 1
        for sentence in re.split(_SENTENCE_END, text):
            words = re.findall(r"[A-Za-z][A-Za-z0-9\-]*", sentence)
            for word in words:
                spellings[word] += 1
                if re.search(r"[a-z][A-Z]", word):
                    marked.add(word.lower())
            if len(words) < 6 or not sentence.rstrip().endswith((".", "!", "?")):
                continue
            for word in words[1:]:
                if len(word) > 1:
                    seen[word.lower()] += 1
                    if word[0].isupper():
                        capitalized[word.lower()] += 1
    marked |= {w for w, n in trademarked.items() if n >= 3}
    proper = {w for w, n in seen.items() if n >= 3 and capitalized[w] / n >= 0.8}
    brands = {}
    for word, _ in spellings.most_common():
        key = word.lower()
        if key in marked | proper and key not in brands:
            brands[key] = word
    return brands


_SENTENCE_END = r"(?<=[.!?।॥])\s+|\n+"


def _sentences(text: str):
    return [s.strip() for s in re.split(_SENTENCE_END, text) if len(s.strip()) >= MIN_SENTENCE_CHARS]


# A Hindi word never starts with a vowel sign, virama or nukta (ि ् ़ ...); PDF text
# extraction that scrambles Devanagari leaves many of them word-initial.
_DEVANAGARI_WORD = re.compile(r"[\u0900-\u097F]+")
_DEPENDENT = re.compile(r"^[\u0900-\u0903\u093A-\u094F\u0951-\u0957\u0962\u0963]")


def _garbled_hindi(text: str) -> bool:
    words = _DEVANAGARI_WORD.findall(text)
    return len(words) >= 30 and sum(bool(_DEPENDENT.match(w)) for w in words) / len(words) > 0.05


def _chunk(text: str):
    text = re.sub(r"[★☆]+", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= CHUNK_CHARS:
        return [text]

    chunks, start = [], 0
    while start < len(text):
        end = min(start + CHUNK_CHARS, len(text))
        if end < len(text):
            cut = text.rfind(" ", start + CHUNK_CHARS // 2, end)
            if cut > start:
                end = cut
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return [c for c in chunks if c]


class KnowledgeBase:
    def __init__(self, root: Path, device: str, use_reranker: bool = True):
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)
        log.info("Loading embedding model %s on %s ...", EMBED_MODEL_ID, device)
        self.device = device
        self.tokenizer = AutoTokenizer.from_pretrained(EMBED_MODEL_ID)
        dtype = torch.float16 if device == "cuda" else torch.float32
        self.model = AutoModel.from_pretrained(EMBED_MODEL_ID, torch_dtype=dtype).to(device).eval()
        self.reranker = self.rerank_tokenizer = None
        if use_reranker:
            try:
                self.rerank_tokenizer = AutoTokenizer.from_pretrained(RERANK_MODEL_ID)
                self.reranker = AutoModelForSequenceClassification.from_pretrained(
                    RERANK_MODEL_ID, torch_dtype=dtype).to(device).eval()
                log.info("Reranker %s loaded.", RERANK_MODEL_ID)
            except Exception as e:  # not downloaded: bge-m3's order is used
                self.reranker = self.rerank_tokenizer = None
                log.warning("Reranker not available (%s) - passages keep the search order.", e)
        self._indexes = {}
        self._warned_legacy = set()
        self._lock = threading.Lock()

    def _folder(self, avatar: str) -> Path:
        if not AVATAR_NAME.match(avatar):
            raise ValueError(f"invalid knowledge folder name: {avatar!r}")
        return self.root / avatar

    def _embed(self, texts):
        # bge-m3 dense embedding = normalized CLS token (same as its reference implementation).
        out = []
        for i in range(0, len(texts), 16):
            batch = self.tokenizer(
                texts[i : i + 16], padding=True, truncation=True, max_length=1024, return_tensors="pt"
            ).to(self.device)
            with torch.no_grad():
                cls = self.model(**batch).last_hidden_state[:, 0]
            out.append(torch.nn.functional.normalize(cls.float(), dim=-1).cpu().numpy())
        return np.vstack(out).astype(np.float32)

    def _scan(self, folder: Path):
        files = []
        if not folder.is_dir():
            return files
        for path in sorted(folder.iterdir()):
            if not path.is_file():
                continue
            ext = path.suffix.lower()
            if ext in SUPPORTED:
                files.append(path)
            elif ext in LEGACY and path not in self._warned_legacy:
                self._warned_legacy.add(path)
                log.warning("Skipping %s - old binary format; re-save it as .docx/.pptx.", path.name)
        return files

    def _file_entry(self, path: Path, cache_dir: Path):
        st = path.stat()
        key = hashlib.sha1(
            f"{path.name}|{st.st_size}|{st.st_mtime_ns}|{EMBED_MODEL_ID}|{CHUNK_CHARS}|{CHUNK_OVERLAP}|{INDEX_FORMAT}".encode()
        ).hexdigest()
        cache_file = cache_dir / f"{key}.npz"

        if cache_file.exists():
            # Close promptly - an open handle blocks deleting it later on Windows.
            with np.load(cache_file) as data:
                return (key, list(data["texts"]), list(data["sources"]), data["emb"].astype(np.float32),
                        data["sent_emb"].astype(np.float32), data["sent_parent"])

        log.info("Indexing %s ...", path.name)
        texts, sources = [], []
        try:
            for label, unit_text in _merge_small_units(_extract_units(path)):
                for chunk in _chunk(unit_text):
                    texts.append(chunk)
                    sources.append(label)
        except Exception:
            log.exception("Could not read %s - skipping it.", path.name)
            return key, [], [], np.zeros((0, 0), np.float32), np.zeros((0, 0), np.float32), np.zeros(0, np.int32)

        if _garbled_hindi("\n".join(texts)):
            log.warning("The Hindi text in %s looks scrambled (common with Hindi PDFs) - answers from it may be wrong. "
                        "Save it as .docx or .txt instead.", path.name)

        sentences, parents = [], []
        for i, t in enumerate(texts):
            for s in _sentences(t):
                sentences.append(s)
                parents.append(i)
        if not texts:
            log.warning("%s contains no extractable text (scanned images aren't supported).", path.name)
            emb = sent_emb = np.zeros((0, 0), np.float32)
        else:
            # The title is searched but never shown to the LLM, so it can't end up
            # in a spoken reply ("Altscape Mobile Deck" helps a mobile question
            # find that deck).
            title = _title(path)
            emb = self._embed([f"{title}\n{t}" for t in texts])
            sent_emb = self._embed([f"{title}\n{s}" for s in sentences]) if sentences else np.zeros((0, emb.shape[1]), np.float32)
        sent_parent = np.array(parents, np.int32)
        np.savez(cache_file, texts=np.array(texts), sources=np.array(sources), emb=emb.astype(np.float16),
                 sent_emb=sent_emb.astype(np.float16), sent_parent=sent_parent)
        log.info("Indexed %s: %d passages, %d sentences.", path.name, len(texts), len(sentences))
        return key, texts, sources, emb, sent_emb, sent_parent

    def _load_or_build(self, avatar: str) -> _Index:
        folder = self._folder(avatar)
        files = self._scan(folder)
        signature = tuple((p.name, p.stat().st_size, p.stat().st_mtime_ns) for p in files)

        cached = self._indexes.get(avatar)
        if cached is not None and cached.signature == signature:
            return cached

        cache_dir = folder / INDEX_DIRNAME
        cache_dir.mkdir(parents=True, exist_ok=True)

        texts, sources, embs, sent_embs, sent_parents, live_keys = [], [], [], [], [], set()
        for path in files:
            key, t, s, e, se, sp = self._file_entry(path, cache_dir)
            live_keys.add(key)
            if t:
                if len(sp):
                    sent_embs.append(se)
                    sent_parents.append(sp + len(texts))  # file-local passage numbers -> global
                texts.extend(t)
                sources.extend(s)
                embs.append(e)

        for stale in cache_dir.glob("*.npz"):
            if stale.stem not in live_keys:
                try:
                    stale.unlink()
                except OSError:
                    log.warning("Could not delete stale cache file %s - will retry next time.", stale.name)

        emb = np.vstack(embs) if embs else np.zeros((0, 0), np.float32)
        brands = _brand_terms(texts)
        sent_emb = np.vstack(sent_embs) if sent_embs else np.zeros((0, 0), np.float32)
        sent_parent = np.concatenate(sent_parents) if sent_parents else np.zeros(0, np.int32)
        index = _Index(signature, texts, sources, emb, brands, sent_emb, sent_parent)
        log.info("Names kept in English for '%s': %s", avatar, ", ".join(sorted(brands.values())) or "none")
        self._indexes[avatar] = index
        if not texts:
            log.warning("Knowledge folder %s has no usable documents - every question will be refused.", folder)
        else:
            log.info("Knowledge '%s' ready: %d documents, %d passages.", avatar, len(files), len(texts))
        return index

    def ingest_all(self):
        for folder in sorted(self.root.iterdir()):
            if folder.is_dir() and AVATAR_NAME.match(folder.name):
                with self._lock:
                    self._load_or_build(folder.name)

    def vocabulary(self, avatar: str, limit: int):
        """Terms to prime speech recognition with, so product/brand names aren't
        mis-heard. Document titles come first - they almost always name the
        company or product - then camel-case words and words containing digits
        from the text (AltCore, iPhone, X1, 5G). Capitalization alone isn't used:
        slide decks are full of Title Case and ALL CAPS headings."""
        with self._lock:
            index = self._load_or_build(avatar)
        folder = self._folder(avatar)

        seen, words = set(), []

        def add(word):
            # Not record codes (CL03, SV01), ordinals (1st, 90th) or registration numbers - the P&I document's filled
            # the hint with them (2026-10-08). Real names with digits stay (G20, MotoGP, 5G).
            if re.fullmatch(r"[A-Z]{1,4}0\d|\d+(?:st|nd|rd|th)|\w{13,}", word):
                return
            if len(word) > 1 and word.lower() not in seen and word.lower() not in _STOPWORDS:
                seen.add(word.lower())
                words.append(word)

        for path in self._scan(folder):
            for word in re.findall(r"[A-Za-z][A-Za-z0-9]*", path.stem):
                add(word)

        counts = Counter(
            word
            for text in index.texts
            for word in re.findall(r"\b[A-Za-z]*[a-z][A-Z][A-Za-z0-9]*\b|\b[A-Za-z]+\d[A-Za-z0-9]*\b|\b\d+[A-Za-z]+\b", text)
        )
        for word, _ in counts.most_common():
            add(word)
        return words[:limit]

    def topic_names(self, avatar: str, limit: int = 3):
        """The brands/products an avatar is about, for avatars whose profile doesn't
        list them: names the documents treat as names (see _brand_terms) ranked by
        how many document titles carry them, then by how often they're mentioned.
        Short all-caps acronyms (AI, PDF) are skipped. Spelled the way the documents
        usually write it, capitalized if they ever do (altcore -> Altcore)."""
        with self._lock:
            index = self._load_or_build(avatar)
        titles = [_title(p).lower() for p in self._scan(self._folder(avatar))]
        spellings = Counter(w for t in index.texts for w in re.findall(r"[A-Za-z][A-Za-z0-9\-]*", t))
        mentions = Counter()
        for word, n in spellings.items():
            mentions[word.lower()] += n

        ranked = []
        for key, spelling in index.brands.items():
            if len(spelling) <= 2 or (spelling.isupper() and len(spelling) <= 4):
                continue
            in_titles = sum(bool(re.search(rf"\b{re.escape(key)}\b", t)) for t in titles)
            ranked.append((in_titles, mentions[key], key))
        ranked.sort(reverse=True)
        # Names from the titles if any; otherwise the most-mentioned ones.
        if ranked and ranked[0][0] > 0:
            ranked = [r for r in ranked if r[0] > 0]

        names = []
        for _, _, key in ranked[:limit]:
            forms = [(n, w) for w, n in spellings.items() if w.lower() == key]
            capitalized = [f for f in forms if f[1][0].isupper() and not f[1].isupper()]
            name = max(capitalized or forms)[1]
            names.append(name if capitalized else name[0].upper() + name[1:].lower() if name.isupper() else name[0].upper() + name[1:])
        return names

    def brand_terms(self, avatar: str) -> dict:
        with self._lock:
            return dict(self._load_or_build(avatar).brands)

    def _rerank(self, question, passages):
        """Reranker scores (0-1) of each passage as an answer to the question."""
        scores = []
        for i in range(0, len(passages), 8):
            batch = self.rerank_tokenizer(
                [[question, p] for p in passages[i : i + 8]], padding=True, truncation=True, max_length=1024,
                return_tensors="pt").to(self.device)
            with torch.no_grad():
                scores.append(torch.sigmoid(self.reranker(**batch).logits.view(-1).float()).cpu().numpy())
        return np.concatenate(scores)

    def retrieve(self, avatar: str, text: str, top_k: int, rerank: bool = True, rank_text: str = ""):
        """Returns (best passage score, best single-sentence score, passages). The
        passage score is what the relevance gate uses; sentence scores run higher
        for everything, junk included, so they only rank passages and back up
        borderline questions. rerank: reorder the closest passages with the reranker
        (if loaded) - the two gate scores are the same either way. rank_text: the question as it should be searched for
        ranking only (the P&I bots search 'the company' instead of 'P&I', which the models don't read as Pavilions &
        Interiors: right passage among 6 in 25/26 vs 17/26); the gate scores stay the question's own."""
        with self._lock:
            index = self._load_or_build(avatar)
            if not index.texts:
                return 0.0, 0.0, []
            # Speech recognition writes brand names in a Hindi question in Devanagari
            # (अल्टस्केप); put them back in the documents' spelling so they match.
            if _DEVANAGARI.search(text) and index.brands:
                text, _ = restore_english(" ".join(index.brands.values()), text)
            query = self._embed([text])[0]
            if rank_text and rank_text != text:
                text = rank_text
                rank_query = self._embed([rank_text])[0]
            else:
                rank_query = query

        def match(q):
            whole = index.emb @ q
            sentence = np.full(len(index.texts), -1.0, np.float32)
            if len(index.sent_parent):
                np.maximum.at(sentence, index.sent_parent, index.sent_emb @ q)
            return whole, sentence

        scores, best_sentence = match(query)
        rank_scores, rank_sentence = (scores, best_sentence) if rank_query is query else match(rank_query)
        # A passage ranks by its whole text or its best sentence, whichever matches better.
        order = np.argsort(-np.maximum(rank_scores, rank_sentence))
        reranked = None
        if rerank and self.reranker is not None and len(order) > 1:
            candidates = order[:RERANK_CANDIDATES]
            with self._lock:
                reranked = self._rerank(text, [index.texts[i] for i in candidates])
            order = candidates[np.argsort(-reranked, kind="stable")]
            reranked = dict(zip(candidates.tolist(), reranked.tolist()))
        order = order[: max(top_k, 1)]
        chunks = [
            {"text": index.texts[i], "source": index.sources[i], "score": float(scores[i])}
            | ({"rerank": reranked[int(i)]} if reranked else {})
            for i in order[:top_k]
        ]
        return float(scores.max()), float(best_sentence.max()), chunks
